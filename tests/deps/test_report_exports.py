"""The audit report as PDF, Word and markdown — one block list, three renderings.

The first exports were a printout of the whole interactive page (58 pages) and
a .docx of typed bullets and stray markdown markers. These pin what makes the
exported files a report: a real PDF with page numbers and Cyrillic, a Word file
with tables and styles and no raw markdown, and a markdown that stays valid.
Fixtures are generic (`acme-api`, `shop-web`); no real repository appears here.
"""

from __future__ import annotations

import io
import re
from types import SimpleNamespace

import pytest

from src.deps.document import build_markdown, build_report, footer_label
from src.docs.blocks import (
    Bullets,
    Heading,
    Table,
    parse_inline,
    parse_markdown,
    to_markdown_text,
)
from src.docs.docx_blocks import blocks_to_docx
from src.docs.pdf_blocks import blocks_to_pdf

UA = "Звіт про залежності: оновити бібліотеки"


def _f(repo, pkg, cur, latest, out="minor", sev="none", vulns=None, src="lock",
       eco="PyPI", dev=False):
    return SimpleNamespace(
        repo_slug=repo, ecosystem=eco, package=pkg, current_version=cur,
        latest_version=latest, outdated=out, is_dev=dev, vulns=vulns or [],
        severity=sev, recommendation="ok", version_source=src)


def _run():
    return SimpleNamespace(id="run-sample", summary={
        "repos_scanned": 3, "repos_total": 4, "packages": 40, "unique_packages": 30,
        "transitive": 10, "outdated": 3, "vulnerable": 3,
        "by_severity": {"critical": 1, "high": 1, "low": 1},
        "groups": {"update_now": 2, "update_safe": 1, "plan_major": 1},
        "sources": {"pip-audit": 2, "osv": 1},
        "coverage": {"checked": 3, "covered_by_osv": 1, "partial": 1},
        "version_provenance": {"lock": 30, "pin": 5, "range": 5},
        "dropped_inapplicable": {"total": 4, "by_source": {"pip-audit": 4}},
        "range_hygiene": {"total": 1, "items": [
            {"repo": "shop-web", "package": "ui-kit", "spec": "^8",
             "advisories": 1, "worst": "high"}]},
        "npm_chains": {"total": 1, "items": [
            {"repo": "shop-web", "package": "build-tool", "pulls_in": ["shell-quote"],
             "fix_via": {"shell-quote": "1.8.3"}, "severity": "critical"}]},
        "not_checked": [{"repo": "acme-api", "ecosystem": "PyPI", "tool": "pip-audit",
                         "status": "partial", "reason": "no lock file"}],
        "ai_report_engine": "claude_code",
        "ai_report": f"# {UA}\n\n## Висновок\n\nРизик **високий**, `django` до 4.2.27.\n\n"
                     "- перше\n- друге\n\n| Крок | Хто |\n|---|---|\n| Оновити | backend |\n",
    })


def _rows():
    return [
        _f("acme-api", "django", "4.2.25", "6.1", "major", "critical",
           [{"cve": "CVE-2026-0001", "fixed_in": "4.2.27"}]),
        _f("shop-web", "django", "4.2.25", "6.1", "major", "high",
           [{"cve": "CVE-2026-0001", "fixed_in": "4.2.27"}]),
        _f("acme-api", "ecdsa", "0.19.0", "0.19.1", "patch", "low",
           [{"cve": "CVE-2026-0003"}], src="pin"),
        _f("shop-web", "ui-kit", "8.0", "9.1", "major", src="range", eco="npm"),
        _f("acme-worker", "requests", "2.32.0", "2.33.0", "patch"),
    ]


@pytest.fixture(scope="module")
def blocks():
    return build_report(_run(), _rows(), generated_at="2026-10-30", workspace="Acme")


# ─── PDF ─────────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def pdf_pages(blocks):
    from pypdf import PdfReader

    data = blocks_to_pdf(blocks, footer=footer_label("2026-10-30"))
    assert data.startswith(b"%PDF")
    reader = PdfReader(io.BytesIO(data))
    return data, [p.extract_text() for p in reader.pages], reader


def test_the_pdf_is_a_pdf_with_page_numbers(pdf_pages):
    _, pages, reader = pdf_pages
    assert len(pages) >= 2
    assert "Page 1 of" in pages[0]
    assert f"Page {len(pages)} of {len(pages)}" in pages[-1]
    assert reader.metadata.title == "Celmis — Dependency Audit"


def test_the_pdf_holds_cyrillic_a_table_row_and_the_title_block(pdf_pages):
    _, pages, _ = pdf_pages
    text = "\n".join(pages)
    # Embedded font, not a system one: Cyrillic must survive extraction.
    assert UA in text.replace("\n", " ")
    assert "Workspace" in text and "Acme" in text and "run-sample" in text
    # A findings-table row keeps its cells together on one line.
    assert re.search(r"ecdsa\s+0\.19\.0\s+pin\s+no fix yet", text)
    assert "Findings by repository" in text


def test_the_pdf_embeds_its_fonts(pdf_pages):
    data, _, reader = pdf_pages
    fonts = {str(f) for page in reader.pages
             for f in page["/Resources"]["/Font"].values()}
    assert fonts, "no fonts at all"
    # Every font resource is a subset embedded as a file, none a bare name.
    assert b"/FontFile2" in data
    assert b"NotoSans" in data


def test_a_long_report_does_not_pile_up_into_noise():
    """250 outdated rows is the appendix cap; 600 outdated packages must not
    make a 600-row table."""
    rows = [_f("acme-api", f"pkg{i}", "1.0", "2.0", "major") for i in range(600)]
    blocks = build_report(_run(), rows, generated_at="2026-10-30")
    data = blocks_to_pdf(blocks)
    from pypdf import PdfReader
    assert len(PdfReader(io.BytesIO(data)).pages) < 20


# ─── DOCX ────────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def docx(blocks):
    import docx as python_docx

    data = blocks_to_docx(blocks, footer=footer_label("2026-10-30"))
    return data, python_docx.Document(io.BytesIO(data))


def test_the_docx_has_real_tables_and_styles(docx):
    _, doc = docx
    assert len(doc.tables) >= 5
    styles = {p.style.name for p in doc.paragraphs}
    assert {"Title", "Heading 2", "Heading 3"} <= styles
    findings = next(t for t in doc.tables if t.rows[0].cells[0].text == "Package"
                    and "Fixed in" in [c.text for c in t.rows[0].cells]
                    and any(r.cells[0].text == "ecdsa" for r in t.rows))
    row = next(r for r in findings.rows if r.cells[0].text == "ecdsa")
    cells = [c.text for c in row.cells]
    assert "0.19.0" in cells and "no fix yet" in cells and "LOW" in cells


def test_the_docx_has_no_raw_markdown_or_typed_bullets(docx):
    _, doc = docx
    texts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            texts += [c.text for c in r.cells]
    joined = "\n".join(texts)
    assert "**" not in joined and "`" not in joined
    assert "•" not in joined          # bullets are list numbering, not text
    assert not re.search(r"(^|\s)_\w", joined)
    # Cyrillic from the AI report is real text, and its table became a table.
    assert UA in joined
    assert any(t.rows[0].cells[0].text == "Крок" for t in doc.tables)


def test_the_docx_lists_are_lists_and_the_footer_numbers_pages(docx):
    import zipfile

    data, doc = docx
    assert any(p.style.name == "List Bullet" for p in doc.paragraphs)
    z = zipfile.ZipFile(io.BytesIO(data))
    footer = z.read("word/footer1.xml").decode()
    assert "PAGE" in footer and "NUMPAGES" in footer
    # Header row repeats across pages; severity cell carries a fill.
    document = z.read("word/document.xml").decode()
    assert "<w:tblHeader/>" in document
    assert 'w:fill="B91C1C"' in document


def test_the_docx_shows_what_the_accuracy_work_introduced(docx):
    _, doc = docx
    text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "≥ 8.0" in text                    # range row shows its floor
    assert "Only a range" in text             # version_provenance
    assert "Audited by its native tool" in text   # coverage
    assert "advisories were not counted" in "\n".join(p.text for p in doc.paragraphs)
    assert "shell-quote" in text              # npm chain
    assert "^8" in text                       # range hygiene


# ─── markdown ────────────────────────────────────────────────────────

def test_the_markdown_is_valid_and_ordered():
    md = build_markdown(_run(), _rows(), generated_at="2026-10-30")
    order = ["## Summary", "## What to do now", "## Coverage & sources",
             "## Findings by repository", "## Appendix A", "## Appendix B"]
    positions = [md.index(h) for h in order]
    assert positions == sorted(positions)
    # Every pipe table is rectangular, so it renders on any markdown viewer.
    tables = re.findall(r"((?:^\|.*\|\n)+)", md, re.M)
    assert tables
    for t in tables:
        widths = {line.count("|") - line.count("\\|") for line in t.splitlines()}
        assert len(widths) == 1, t
    assert "**" in md                         # emphasis is markdown here, by design
    assert md.endswith("\n") and not md.endswith("\n\n")


def test_vulnerable_packages_come_before_the_outdated_appendix():
    md = build_markdown(_run(), _rows(), generated_at="2026-10-30")
    assert md.index("| ecdsa |") < md.index("## Appendix A")
    # The clean outdated package is in the appendix table, not under a repo.
    assert md.index("| acme-worker | requests") > md.index("## Appendix A")


def test_every_row_says_where_its_version_came_from(blocks):
    tables = [b for b in blocks if isinstance(b, Table) and "Src" in b.header]
    assert tables
    sources = {row[tables[0].header.index("Src")] for t in tables for row in t.rows}
    assert sources <= {"pin", "lock", "resolved", "range", "reported", ""}
    assert {"lock", "pin"} <= sources


def test_a_run_from_before_version_provenance_has_no_empty_src_column():
    rows = _rows()
    for r in rows:
        r.version_source = None
    old = build_report(_run(), rows, generated_at="2026-10-30")
    tables = [b for b in old if isinstance(b, Table) and "Severity" in b.header
              and "Fixed in" in b.header and "Latest" in b.header]
    assert tables and all("Src" not in t.header for t in tables)
    assert all(len(row) == len(t.header) for t in tables for row in t.rows)
    assert blocks_to_pdf(old).startswith(b"%PDF")


# ─── the converter itself ────────────────────────────────────────────

def test_inline_markup_becomes_runs_not_text():
    runs = parse_inline("a **b** `c` _d_ snake_case_name 2*3")
    assert [(r.text, r.bold, r.code, r.italic) for r in runs if r.bold or r.code or r.italic] \
        == [("b", True, False, False), ("c", False, True, False), ("d", False, False, True)]
    assert "snake_case_name" in "".join(r.text for r in runs)


def test_markdown_round_trips_through_blocks():
    md = "# T\n\n## H\n\ntext\n\n- a\n- b\n\n1. x\n2. y\n\n| A | B |\n| --- | --- |\n| 1 | 2 |\n"
    blocks = parse_markdown(md)
    assert isinstance(blocks[0], Heading) and blocks[0].level == 1
    assert [type(b).__name__ for b in blocks] == [
        "Heading", "Heading", "Para", "Bullets", "Bullets", "Table"]
    assert isinstance(blocks[3], Bullets) and not blocks[3].ordered and blocks[4].ordered
    assert parse_markdown(to_markdown_text(blocks)) == blocks
